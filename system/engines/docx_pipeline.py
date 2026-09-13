import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from system.core.config import ROOT_DIR

class DocxPipeline:
    def __init__(self):
        pass

    def export_chapter_to_docx(self, title: str, chapter_num: int, content_md: str, output_docx_path: str):
        os.makedirs(os.path.dirname(output_docx_path), exist_ok=True)
        doc = docx.Document()

        # Page setup - Margins
        sections = doc.sections
        for s in sections:
            s.top_margin = Inches(1)
            s.bottom_margin = Inches(1)
            s.left_margin = Inches(1.2)
            s.right_margin = Inches(1)

        # Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(f"PH? TR?I\n{title.upper()}")
        run_title.font.name = "Times New Roman"
        run_title.font.size = Pt(18)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(30, 30, 30)

        # Meta
        p_meta = doc.add_paragraph()
        p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_meta = p_meta.add_run("T?c gi?: Author & Novel OS Co-Author Engine | B?n Th?o Ch?nh Th?c")
        run_meta.font.name = "Times New Roman"
        run_meta.font.size = Pt(10)
        run_meta.font.italic = True
        run_meta.font.color.rgb = RGBColor(100, 100, 100)

        doc.add_paragraph() # Spacer

        # Content lines
        lines = content_md.split("\n")
        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue
            if line_str.startswith("#"):
                continue # B? qua header markdown tr?ng l?p
            
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.first_line_indent = Inches(0.3)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            run = p.add_run(line_str)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(20, 20, 20)

        doc.save(output_docx_path)
        return output_docx_path
